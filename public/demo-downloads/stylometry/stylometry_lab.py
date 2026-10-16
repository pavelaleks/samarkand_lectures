#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ПОЛНАЯ ЛАБОРАТОРНАЯ · Стилиометрия (Burrows' Delta) · Бунин + Чехов

Новичкам в PyCharm:
  1) Откройте этот файл
  2) Run (параметры не нужны)
  3) Смотрите out/ — Excel, CSV, HTML, графики

Natasha здесь НЕ нужна — NER — лекция 6.

Выход:
  • средние Delta: внутри Чехова / внутри Бунина / между авторами
  • матрица Delta, bar charts, MDS-карта
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from stylometry_lib import (  # noqa: E402
    burrows_delta,
    corpus_stats,
    esc,
    fw_vector,
    load_corpus,
    load_meta,
    mds_from_deltas,
    top_fw_diffs,
)

OUT = ROOT / "out"
TEXTS_BUNIN = ROOT / "sample_texts"
TEXTS_CHEKHOV = ROOT / "sample_texts_chekhov"
META_BUNIN = ROOT / "stories_meta.json"
META_CHEKHOV = ROOT / "chekhov_meta.json"

SHOWCASE_PAIRS = [
    ("dama-s-sobachkoj", "ionyich", "Чехов ↔ Чехов · любовные новеллы"),
    ("kryzhovnik", "chelovek-v-futlyare", "Чехов ↔ Чехов · «маленькая трилогия»"),
    ("holodnaya-osen", "kavkaz", "Бунин ↔ Бунин · короткие"),
    ("chistyj-ponedelnik", "natali", "Бунин ↔ Бунин · длинные"),
    ("dama-s-sobachkoj", "holodnaya-osen", "Чехов ↔ Бунин · короткая любовная проза"),
    ("ionyich", "natali", "Чехов ↔ Бунин · провинция / судьба"),
]


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore", delimiter=";")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def write_excel(path: Path, sheets: dict[str, tuple[list[str], list[dict]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, (fields, rows) in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        ws.append(fields)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c, "") for c in fields])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def plot_fw_bars(slug_a: str, slug_b: str, rows, title: str, path: Path) -> None:
    import matplotlib.pyplot as plt

    words = [r[0] for r in rows[:10]]
    fa = [r[1] * 100 for r in rows[:10]]
    fb = [r[2] * 100 for r in rows[:10]]
    y = range(len(words))
    h = 0.35
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.barh([i - h / 2 for i in y], fa, height=h, label=slug_a, color="#0d9488")
    ax.barh([i + h / 2 for i in y], fb, height=h, label=slug_b, color="#6366f1")
    ax.set_yticks(list(y))
    ax.set_yticklabels(words)
    ax.set_xlabel("Доля токенов, %")
    ax.set_title(title)
    ax.legend(loc="lower right", fontsize=8)
    fig.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=120)
    plt.close(fig)


def plot_mds(coords: dict[str, tuple[float, float]], authors: dict[str, str], meta: dict, path: Path) -> None:
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(9, 7))
    for slug, (x, y) in coords.items():
        title = meta.get(slug, {}).get("title", slug)
        is_chekhov = authors.get(slug) == "chekhov"
        color = "#f43f5e" if is_chekhov else "#0d9488"
        ax.scatter(x, y, c=color, s=40, zorder=3)
        ax.annotate(title, (x, y), fontsize=6, alpha=0.85, xytext=(3, 3), textcoords="offset points")
    ax.set_title("MDS по Delta · бирюзовый Бунин · красный Чехов")
    ax.set_xlabel("ось 1")
    ax.set_ylabel("ось 2")
    ax.grid(True, alpha=0.3)
    fig.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=120)
    plt.close(fig)


def write_html(path: Path, pair_rows: list[dict], findings: list[str], summary: dict) -> None:
    pair_tr = "".join(
        f"<tr><td>{esc(r['a_title'])}</td><td>{esc(r['b_title'])}</td>"
        f"<td>{r['delta']:.4f}</td><td>{esc(r['label'])}</td></tr>"
        for r in pair_rows
    )
    findings_li = "".join(f"<li>{esc(x)}</li>" for x in findings)
    body = f"""<!DOCTYPE html><html lang="ru"><head><meta charset="utf-8"/>
<title>Стилиометрия · Бунин и Чехов</title>
<style>body{{font-family:system-ui,sans-serif;max-width:900px;margin:2rem auto;padding:0 1rem}}
table{{border-collapse:collapse;width:100%}} th,td{{border:1px solid #ccc;padding:6px;font-size:13px}}
th{{background:#f0fdfa}} .big{{font-size:1.4rem;font-weight:700}}</style></head><body>
<h1>Учебная стилиометрия (Burrows' Delta)</h1>
<p><strong>Не атрибуция авторства.</strong> Сравниваем служебные профили рассказов Бунина и Чехова.</p>
<h2>Три средних</h2>
<ul>
<li>Внутри Чехова: <span class="big">{summary['mean_chekhov']:.3f}</span></li>
<li>Внутри Бунина: <span class="big">{summary['mean_bunin']:.3f}</span></li>
<li>Чехов ↔ Бунин: <span class="big">{summary['mean_cross']:.3f}</span></li>
</ul>
<h2>Витринные пары</h2>
<table><tr><th>A</th><th>B</th><th>Delta</th><th>Комментарий</th></tr>{pair_tr}</table>
<h2>Выводы для семинара</h2><ul>{findings_li}</ul>
<p>Natasha / NER — только лекция 6.</p>
</body></html>"""
    path.write_text(body, encoding="utf-8")


def main() -> None:
    bunin = load_corpus(TEXTS_BUNIN)
    chekhov = load_corpus(TEXTS_CHEKHOV) if TEXTS_CHEKHOV.exists() else {}
    if not chekhov:
        print("WARNING: нет sample_texts_chekhov/ — будет только Бунин")

    meta = {}
    meta.update(load_meta(META_BUNIN))
    meta.update(load_meta(META_CHEKHOV))

    authors: dict[str, str] = {}
    corpus: dict[str, str] = {}
    for slug, text in bunin.items():
        corpus[slug] = text
        authors[slug] = "bunin"
    for slug, text in chekhov.items():
        corpus[slug] = text
        authors[slug] = "chekhov"
        if slug not in meta:
            meta[slug] = {"title": slug, "slug": slug}

    slugs = sorted(corpus.keys())
    vectors: dict[str, list[float]] = {}
    tokens: dict[str, int] = {}
    for slug, text in corpus.items():
        v, n = fw_vector(text)
        vectors[slug] = v
        tokens[slug] = n

    means, stds = corpus_stats([vectors[s] for s in slugs])

    matrix: dict[tuple[str, str], float] = {}
    matrix_rows: list[dict] = []
    for i, a in enumerate(slugs):
        for b in slugs[i + 1 :]:
            d = burrows_delta(vectors[a], vectors[b], means, stds)
            matrix[(a, b)] = d
            ta = meta.get(a, {}).get("title", a)
            tb = meta.get(b, {}).get("title", b)
            matrix_rows.append(
                {
                    "a": ta,
                    "b": tb,
                    "slug_a": a,
                    "slug_b": b,
                    "author_a": authors[a],
                    "author_b": authors[b],
                    "delta": round(d, 4),
                }
            )
    matrix_rows.sort(key=lambda r: r["delta"])

    def mean_of(pred):
        vals = [r["delta"] for r in matrix_rows if pred(r)]
        return sum(vals) / len(vals) if vals else 0.0

    mean_bunin = mean_of(lambda r: r["author_a"] == r["author_b"] == "bunin")
    mean_chekhov = mean_of(lambda r: r["author_a"] == r["author_b"] == "chekhov")
    mean_cross = mean_of(lambda r: r["author_a"] != r["author_b"])

    pair_report: list[dict] = []
    OUT.mkdir(parents=True, exist_ok=True)
    plots = OUT / "plots"
    plots.mkdir(exist_ok=True)

    for sa, sb, label in SHOWCASE_PAIRS:
        if sa not in vectors or sb not in vectors:
            continue
        d = burrows_delta(vectors[sa], vectors[sb], means, stds)
        ta = meta.get(sa, {}).get("title", sa)
        tb = meta.get(sb, {}).get("title", sb)
        pair_report.append(
            {"a_title": ta, "b_title": tb, "delta": d, "label": label, "slug_a": sa, "slug_b": sb}
        )
        diffs = top_fw_diffs(vectors[sa], vectors[sb])
        plot_fw_bars(
            sa,
            sb,
            diffs,
            f"|Δ| служебных: {ta} vs {tb} (Delta={d:.3f})",
            plots / f"{sa}__{sb}_fw.png",
        )

    coords = mds_from_deltas(slugs, matrix)
    plot_mds(coords, authors, meta, plots / "mds_authors.png")

    metrics_rows = []
    for slug in slugs:
        m = meta.get(slug, {})
        metrics_rows.append(
            {
                "slug": slug,
                "title": m.get("title", slug),
                "author": authors[slug],
                "tokens": tokens[slug],
                "fw_share_pct": round(sum(vectors[slug]) * 100, 2),
                "mds_x": round(coords.get(slug, (0, 0))[0], 4),
                "mds_y": round(coords.get(slug, (0, 0))[1], 4),
            }
        )

    summary = {
        "mean_bunin": mean_bunin,
        "mean_chekhov": mean_chekhov,
        "mean_cross": mean_cross,
    }
    findings = [
        f"Средняя Delta внутри Чехова ≈ {mean_chekhov:.3f} — рассказы сходны по служебному профилю.",
        f"Средняя Delta внутри Бунина ≈ {mean_bunin:.3f} — единство цикла «Тёмные аллеи».",
        f"Средняя Delta Чехов ↔ Бунин ≈ {mean_cross:.3f} — типично дальше внутриавторских средних.",
        "Сравнивайте тексты похожей длины. Delta ≠ судебная атрибуция.",
    ]

    write_csv(OUT / "delta_matrix.csv", ["a", "b", "slug_a", "slug_b", "author_a", "author_b", "delta"], matrix_rows)
    write_csv(
        OUT / "metrics.csv",
        ["slug", "title", "author", "tokens", "fw_share_pct", "mds_x", "mds_y"],
        metrics_rows,
    )

    try:
        write_excel(
            OUT / "stylometry_bunin.xlsx",
            {
                "summary": (
                    ["metric", "value"],
                    [
                        {"metric": "mean_delta_within_chekhov", "value": round(mean_chekhov, 4)},
                        {"metric": "mean_delta_within_bunin", "value": round(mean_bunin, 4)},
                        {"metric": "mean_delta_chekhov_vs_bunin", "value": round(mean_cross, 4)},
                    ],
                ),
                "metrics": (
                    ["slug", "title", "author", "tokens", "fw_share_pct", "mds_x", "mds_y"],
                    metrics_rows,
                ),
                "pairs_showcase": (
                    ["a_title", "b_title", "delta", "label"],
                    [{k: r[k] for k in ("a_title", "b_title", "delta", "label")} for r in pair_report],
                ),
                "delta_matrix": (
                    ["a", "b", "slug_a", "slug_b", "author_a", "author_b", "delta"],
                    matrix_rows,
                ),
            },
        )
    except ImportError:
        print("openpyxl не установлен — Excel пропущен, есть CSV.")

    write_html(OUT / "report.html", pair_report, findings, summary)

    print("=== Стилиометрия · готово ===")
    print(f"Бунин: {sum(1 for a in authors.values() if a=='bunin')} · Чехов: {sum(1 for a in authors.values() if a=='chekhov')}")
    print(f"Средняя Delta внутри Чехова: {mean_chekhov:.4f}")
    print(f"Средняя Delta внутри Бунина: {mean_bunin:.4f}")
    print(f"Средняя Delta Чехов ↔ Бунин: {mean_cross:.4f}")
    for r in pair_report:
        print(f"  {r['a_title']} ↔ {r['b_title']}: Delta={r['delta']:.4f}")
    print(f"\nФайлы в {OUT}/")


if __name__ == "__main__":
    main()
